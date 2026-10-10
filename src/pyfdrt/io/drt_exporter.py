import io
import os
import numpy as np
import pandas as pd
from typing import Callable, Optional
from ..core.DRTOptimizer import DRTOptimizer
from .circuit_cdc import circuit_cdc


class ExcelExportError(Exception):
    """Raised when :func:`export_results_to_excel` cannot write the workbook.

    Carries the target ``file_path`` and is chained (``raise ... from``) to the
    underlying pandas / openpyxl / OS error, so a caller can report a specific
    failure instead of the exporter swallowing it and returning ``False``.
    See docs/specs/drt-io.md section 5.
    """

    def __init__(self, file_path: str, message: str):
        self.file_path = file_path
        super().__init__(f"Failed to save Excel to '{file_path}': {message}")


def _embed_circuit_diagram(worksheet, results, log, inductance=0.0):
    """Put the circuit schematic under the parameter table (drt-io.md 5.1).

    The picture goes on the same sheet as the numbers it describes, two rows
    below the table, so a reader sees the circuit and its parameters together
    without changing sheet. It carries element names only - the values are the
    rows directly above it, and a second copy in the picture would be a second
    thing to keep in agreement.

    Written last and reported rather than raised on failure: embedding needs
    Pillow, which openpyxl requires for pictures but does not declare, and a
    workbook without a picture is worth far more than no workbook.
    """
    anchor_row = worksheet.max_row + 2
    try:
        import matplotlib.pyplot as plt
        from openpyxl.drawing.image import Image as _XLImage

        from .drt_plotter import _circuit_figure

        fig, _ax = _circuit_figure(results, inductance)
        buffer = io.BytesIO()
        fig.savefig(buffer, format='png', facecolor='white')
        plt.close(fig)
        buffer.seek(0)
        image = _XLImage(buffer)
        image.anchor = f'A{anchor_row}'
        worksheet.add_image(image)
    except Exception as e:                       # pragma: no cover - guarded path
        worksheet[f'A{anchor_row}'] = 'Circuit diagram'
        worksheet[f'B{anchor_row}'] = f'not embedded ({type(e).__name__})'
        log(f"-> [Warning] Circuit diagram not embedded: {e}")


def export_results_to_excel(file_path: str, optimizer: DRTOptimizer, results: dict,
                            logger: Optional[Callable[[str], None]] = None,
                            inductance: float = 0.0):
    """
    Save the DRT analysis and optimization results to an Excel file.

    Returns ``True`` on success. Raises :class:`ExcelExportError` (chained from
    the underlying error) when the workbook cannot be written, after logging
    ``-> [Error] Failed to save Excel: ...``. It no longer reports failure as a
    ``False`` return, so a caller cannot mistake a write failure for success.
    """
    log = logger or print
    try:
        opt = optimizer

        # 1. DRT spectrum data — all arrays on the tau grid (fixed 10 pts/decade).
        drt_cols = {
            'f_Hz':         1.0 / (2.0 * np.pi * opt.tau),
            'omega_rad_s':  1.0 / opt.tau,
            'Tau_s':        opt.tau,
            'Log10_Tau':    opt.log_tau,
            'DRT_G_tau':    opt.polished_drt,
        }

        # Peak fitting results
        if (hasattr(opt, 'refined_ecm_params') and opt.refined_ecm_params
                and opt.refined_ecm_params.get('Peaks')):
            peak_sum = np.zeros_like(opt.log_tau)
            for p in opt.refined_ecm_params['Peaks']:
                curve = DRTOptimizer.peak_curve(opt.log_tau, p)
                drt_cols[f"Peak_{p['Peak']}_G_tau"] = curve
                peak_sum += curve
            drt_cols['Peak_fit_sum_G_tau'] = peak_sum
        df_drt = pd.DataFrame(drt_cols)

        # 2. Simulated impedance data: NNLS and ECM on one common grid / R_inf.
        rec = opt.reconstruct_impedance(freq=opt.target_data[:, 0])
        sim_cols = {
            'f_Hz':          rec['freq'],
            'omega_rad_s':   2.0 * np.pi * rec['freq'],
            'Measured_Z_re': opt.target_data[:, 1],
            'Measured_Z_im': opt.target_data[:, 2],
            'NNLS_Z_re':     rec['nnls_re'],
            'NNLS_Z_im':     rec['nnls_im'],
        }
        if rec['ecm_re'] is not None:
            sim_cols['EC_Z_re'] = rec['ecm_re']
            sim_cols['EC_Z_im'] = rec['ecm_im']
        df_sim = pd.DataFrame(sim_cols)

        # 3. Extracted ECM parameters (R, phi, CPE Q, Tau, C_eff per element)
        if results and 'Peaks' in results and results['Peaks']:
            df_params = pd.DataFrame(results['Peaks'])
            if 'Phi' not in df_params:
                df_params['Phi'] = 1.0
            if 'C_eff' not in df_params:
                df_params['C_eff'] = np.where(df_params['R'] > 0,
                                              df_params['Tau'] / df_params['R'], 0.0)
            df_params['R_inf'] = results.get('R_inf', 0.0)
            # Order the most relevant columns first. Phi (refined, impedance-domain
            # CPE exponent) and Phi_deconv (DRT-domain exponent that sets the
            # plotted curve shape) are kept as adjacent, explicitly distinct
            # columns so the table and the plotted spectrum can be compared
            # directly rather than silently implying different peak widths (task #20).
            lead = [c for c in ['Peak', 'R', 'Phi', 'Phi_deconv', 'CPE_Q', 'Tau',
                                'C_eff', 'Freq', 'R_inf'] if c in df_params.columns]
            df_params = df_params[lead + [c for c in df_params.columns if c not in lead]]
            # Exclude redundant / intermediate columns. Phi_deconv is intentionally
            # NOT dropped here (see above).
            drop_cols = ['Freq', 'R_inf', 'Sigma', 'Amplitude', 'Mu',
                         'R_gauss', 'Tau_gauss']
            df_params = df_params.drop(columns=drop_cols, errors='ignore')
        else:
            df_params = pd.DataFrame([{'R_inf': results.get('R_inf', 0.0) if results else 0.0}])

        # Ensure directory exists
        dir_name = os.path.dirname(file_path)
        if dir_name:
            os.makedirs(dir_name, exist_ok=True)

        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            # R_inf is a spectrum-level scalar, not a per-peak quantity: it is
            # written as a labelled cell above the table (task #33) rather than
            # dropped, so the workbook carries the complete equivalent circuit.
            # The parameters table starts at Excel row 3 (startrow=2).
            df_params.to_excel(writer, sheet_name='Equivalent Circuit Parameters',
                               index=False, startrow=2)
            ws = writer.sheets['Equivalent Circuit Parameters']
            ws['A1'] = 'R_inf (ohm)'
            ws['B1'] = float(results.get('R_inf', 0.0) if results else 0.0)
            # The Warburg-open diffusion element (drt-deconvolution.md 2.4) is
            # likewise spectrum-level: when present it occupies the otherwise
            # blank spacer row 2 as labelled pairs (task #38). The table still
            # starts at row 3, so table consumers see an unchanged layout.
            wb = results.get('Warburg') if results else None
            if wb:
                ws['A2'] = 'Warburg R_W (ohm)'
                ws['B2'] = float(wb['R_W'])
                ws['C2'] = 'Warburg Tau_W (s)'
                ws['D2'] = float(wb['Tau_W'])
                # A degenerate joint fit withholds C_lim (drt-core 5.5): the
                # label is still written and the value cell left empty, so the
                # sheet says the quantity does not exist for this fit rather
                # than carrying an unsupported number (drt-io.md 2, task #98).
                ws['E2'] = 'Warburg C_lim (F)'
                if wb.get('C_lim') is not None:
                    ws['F2'] = float(wb['C_lim'])
                # The composite's slow ZARC is spectrum-level too, and on
                # LIB_data_2 it carries 48 % of the non-ohmic resistance - more
                # than twice the largest tabulated peak - so a workbook without
                # it understates the circuit by half while looking complete. It
                # stays out of the peak table: that table is the DRT-detected
                # relaxation set, while this element is co-fitted on the
                # residual (drt-io.md 2, drt-deconvolution.md 2.4.2, task #100).
                sz = wb.get('slow_zarc')
                if sz:
                    ws['G2'] = 'Warburg slow ZARC R (ohm)'
                    ws['H2'] = float(sz['R'])
                    ws['I2'] = 'Warburg slow ZARC Tau (s)'
                    ws['J2'] = float(sz['Tau'])
                    ws['K2'] = 'Warburg slow ZARC Phi'
                    ws['L2'] = float(sz['Phi'])
                    # C_eff for parity with the peak table, so every element of
                    # the circuit the schematic draws can be read off this sheet.
                    if sz.get('C_eff') is not None:
                        ws['M2'] = 'Warburg slow ZARC C_eff (F)'
                        ws['N2'] = float(sz['C_eff'])
            # The circuit description code is a spectrum-level scalar like
            # R_inf, so it joins row 1 rather than earning a sheet. C1/D1
            # leaves the table's start at row 3 untouched (drt-io.md 5.1).
            ws['C1'] = 'Circuit (CDC)'
            ws['D1'] = circuit_cdc(results, inductance)
            # The inductance the preprocessing adopted. It entered the code and
            # the schematic in task #140's follow-up but had no cell, so the
            # sheet showed a leading `L` with no value behind it (task #143).
            # Written only when one was adopted: a zero is the gates rejecting
            # the estimate, and a `0` here would read as a measured zero.
            if inductance and inductance > 0:
                ws['E1'] = 'L (H)'
                ws['F1'] = float(inductance)
            df_drt.to_excel(writer,    sheet_name='DRT_Spectrum',   index=False)
            df_sim.to_excel(writer,    sheet_name='Simulated_EIS',  index=False)
            _embed_circuit_diagram(ws, results, log, inductance)

        log(f"-> [OK] Excel saved: {file_path}")
        return True
    except Exception as e:
        log(f"-> [Error] Failed to save Excel: {str(e)}")
        raise ExcelExportError(file_path, str(e)) from e
